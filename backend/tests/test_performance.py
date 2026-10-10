"""Management figures: whether the executive asked about a complaint, and
each estimate's rule with worked examples (coverage, First Call Resolution,
repeats, churn risk, CSAT), the per-executive scorecard, the API and the
download."""

import csv
import io
import json
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.ai.complaint.llm_provider import LLMComplaintProvider
from app.ai.complaint.provider import ComplaintDetectionProvider, ComplaintDetectionResult
from app.ai.llm.client import LLMClient, LLMRequest, LLMResponse
from app.ai.question.provider import QuestionSuggestionProvider
from app.ai.sentiment.provider import SentimentAnalysisProvider, SentimentLabel, SentimentResult
from app.api.app_factory import create_app
from app.api.wiring import build_api_services
from app.core.config import Settings
from app.domain.complaint_coverage import ComplaintCoverageStatus
from app.domain.conversation import Conversation
from app.domain.conversation_coverage import ConversationCoverage
from app.domain.escalation import EscalationLevel
from app.domain.user import User, UserRole
from app.domain.utterance import SpeakerRole, Utterance
from app.security.jwt import create_access_token
from app.services.complaint_analysis_service import ComplaintAnalysisService
from app.services.performance import PerformanceService
from app.services.report_export import export_scorecard
from app.services.reporting import (
    ReportCall,
    ReportComplaint,
    ReportError,
    ReportFilters,
    ReportSource,
)

DAY = 86400.0
START = 1_000 * DAY
NOW = START + 60 * DAY


# ---- Has the executive asked about the complaint? ----

class _LLM(LLMClient):
    def __init__(self, answer) -> None:
        self.answer = answer
        self.prompts: list[str] = []

    def complete(self, request: LLMRequest) -> LLMResponse:
        self.prompts.append(request.prompt)
        return LLMResponse(text=json.dumps(self.answer))


def _conversation() -> Conversation:
    call = Conversation(call_id="call-1")
    call.add_utterance(
        Utterance("u1", "The bill is too high and the car is dirty.", SpeakerRole.CUSTOMER, ("en",), 0.0, 1.0)
    )
    return call


def _detection(category: str, **extra) -> dict:
    return {"category": category, "confidence": 0.9, "evidence": "Said so.", **extra}


def test_the_model_is_asked_whether_each_complaint_was_asked_about():
    llm = _LLM([_detection("Cost", probed=True), _detection("Hygiene", probed=False)])

    results = LLMComplaintProvider(llm).detect(_conversation())

    assert '"probed": <true or false>' in llm.prompts[0]
    assert "an apology or a promise alone is not a question" in llm.prompts[0]
    assert [(r.category, r.probed) for r in results] == [("Cost", True), ("Hygiene", False)]


@pytest.mark.parametrize(
    ("value", "probed"),
    [(True, True), ("true", True), ("Yes", True), (False, False), ("false", False),
     (None, False), (1, False), ("maybe", False), ([], False)],
)
def test_a_missing_or_odd_answer_means_not_asked_about_and_keeps_the_detection(value, probed):
    answer = [_detection("Cost") if value is None else _detection("Cost", probed=value)]

    results = LLMComplaintProvider(_LLM(answer)).detect(_conversation())

    assert [(r.category, r.probed) for r in results] == [("Cost", probed)]


class _Detections(ComplaintDetectionProvider):
    def __init__(self, *rounds) -> None:
        self._rounds = list(rounds)

    def detect(self, conversation, learning_context=()):
        return self._rounds.pop(0)


def _result(category: str, probed: bool) -> ComplaintDetectionResult:
    return ComplaintDetectionResult(category, 0.9, "Said so.", probed)


def test_a_complaint_becomes_probed_once_asked_about_and_stays_so():
    service = ComplaintAnalysisService(
        _Detections(
            [_result("Cost", False), _result("Hygiene", False)],
            [_result("Cost", True), _result("Hygiene", False)],
            # A later answer that forgets the question does not undo it.
            [_result("Cost", False), _result("Hygiene", True)],
        )
    )
    call, coverage = _conversation(), ConversationCoverage(call_id="call-1")

    statuses = []
    for _ in range(3):
        coverage = service.analyze(call, coverage)
        statuses.append({c.category: c.status for c in coverage.complaints})

    detected, probed = ComplaintCoverageStatus.DETECTED, ComplaintCoverageStatus.PROBED
    assert statuses == [
        {"Cost": detected, "Hygiene": detected},
        {"Cost": probed, "Hygiene": detected},
        {"Cost": probed, "Hygiene": probed},
    ]


def test_a_complaint_asked_about_when_first_detected_is_probed_at_once():
    service = ComplaintAnalysisService(_Detections([_result("Cost", True)]))

    coverage = service.analyze(_conversation(), ConversationCoverage(call_id="call-1"))

    assert coverage.get("Cost").status is ComplaintCoverageStatus.PROBED


# ---- The figures ----

class _Calls(ReportSource):
    def __init__(self, *calls: ReportCall) -> None:
        self.all = list(calls)

    def calls(self, filters: ReportFilters, limit: int):
        matching = [
            call
            for call in self.all
            if filters.started_from <= call.start_time < filters.started_to
            and filters.executive_user_id in (None, call.executive_user_id)
            and filters.location_id in (None, call.location_id)
        ]
        return tuple(sorted(matching, key=lambda call: call.start_time)[:limit])


def _call(
    call_id: str,
    day: float,
    complaints: dict[str, str] | None = None,
    *,
    probed: tuple[str, ...] = (),
    customer: str | None = None,
    executive: str | None = "asha",
    tone: SentimentLabel | None = SentimentLabel.NEUTRAL,
    escalation: EscalationLevel | None = None,
) -> ReportCall:
    """A call `day` days into the period, with complaints as {category:
    status now}."""
    return ReportCall(
        call_id=call_id,
        start_time=START + day * DAY,
        executive_user_id=executive,
        executive_name=None if executive is None else executive.title(),
        sentiment=tone,
        complaints=tuple(
            ReportComplaint(category, status, None, category in probed)
            for category, status in (complaints or {}).items()
        ),
        customer_key=customer,
        escalation_level=escalation,
    )


def _report(*calls: ReportCall, days: float = 30, now: float = NOW, **filters):
    service = PerformanceService(_Calls(*calls), clock=lambda: now)
    return service.report(ReportFilters(START, START + days * DAY, **filters))


def test_coverage_is_the_share_of_complaints_the_executive_asked_about():
    report = _report(
        _call("a", 1, {"Cost": "detected", "Hygiene": "detected"}, probed=("Cost",)),
        _call("b", 2, {"TAT": "resolved"}, probed=("TAT",)),
        _call("c", 3, {"Cost": "detected"}),
        _call("d", 4),
    )

    overall = report.overall
    assert (overall.calls, overall.complaints, overall.probed_complaints) == (4, 4, 2)
    assert overall.coverage_score == 0.5


def test_without_complaints_there_is_no_coverage_score():
    assert _report(_call("a", 1)).overall.coverage_score is None
    assert _report().overall.calls == 0


def test_first_call_resolution_holds_when_the_customer_does_not_call_back_about_it():
    report = _report(
        # Called again about the same thing three days later: not resolved first time.
        _call("a1", 1, {"Cost": "detected"}, customer="c:1"),
        _call("a2", 4, {"Cost": "detected"}, customer="c:1"),
        # Called again about something else: the first matter was settled.
        _call("b1", 1, {"Cost": "detected"}, customer="c:2"),
        _call("b2", 4, {"Hygiene": "detected"}, customer="c:2"),
        # Called again about it after more than seven days: settled first time.
        _call("c1", 1, {"TAT": "detected"}, customer="c:3"),
        _call("c2", 9, {"TAT": "detected"}, customer="c:3"),
        # Not judged: unknown customer, and a call without complaints.
        _call("d", 1, {"Cost": "detected"}),
        _call("e", 1, customer="c:4"),
    )

    overall = report.overall
    # Judged: a1, a2, b1, b2, c1, c2; only a1 was called back about.
    assert (overall.fcr_calls, overall.fcr_resolved) == (6, 5)
    assert overall.fcr_rate == pytest.approx(5 / 6)


def test_a_call_back_just_after_the_period_still_counts_and_recent_calls_wait():
    calls = (
        _call("a1", 29, {"Cost": "detected"}, customer="c:1"),
        _call("a2", 33, {"Cost": "detected"}, customer="c:1"),  # after the period
        _call("b", 29.5, {"Cost": "detected"}, customer="c:2"),
    )

    settled = _report(*calls).overall
    assert (settled.fcr_calls, settled.fcr_resolved) == (2, 1)

    # Five days after those calls it is too soon to say they were settled.
    waiting = _report(*calls, now=START + 34 * DAY).overall
    assert (waiting.fcr_calls, waiting.fcr_rate) == (0, None)


def test_a_complaint_is_a_repeat_when_the_customer_raised_it_in_the_last_thirty_days():
    report = _report(
        _call("before", -10, {"Cost": "detected"}, customer="c:1"),  # before the period
        _call("a", 5, {"Cost": "detected", "Hygiene": "detected"}, customer="c:1"),
        _call("old", -40, {"TAT": "detected"}, customer="c:2"),  # too long ago
        _call("b", 5, {"TAT": "detected"}, customer="c:2"),
        _call("c", 6, {"Cost": "detected"}),  # unknown customer: not counted
    )

    overall = report.overall
    assert (overall.known_customer_complaints, overall.repeat_complaints) == (3, 1)
    assert overall.repeat_rate == pytest.approx(1 / 3)
    # The earlier call is history only: it is not one of the period's calls.
    assert overall.calls == 3


@pytest.mark.parametrize(
    ("call", "risk"),
    [
        (dict(tone=SentimentLabel.POSITIVE), "low"),
        (dict(complaints={"Cost": "detected"}), "low"),  # 1: open complaint
        (dict(tone=SentimentLabel.NEGATIVE, complaints={"Cost": "detected"}), "medium"),  # 1+1
        (dict(tone=SentimentLabel.NEGATIVE, complaints={"Cost": "resolved"}), "low"),  # 1
        (dict(tone=SentimentLabel.FRUSTRATED), "medium"),  # 2
        (dict(tone=SentimentLabel.ESCALATING), "medium"),  # 3
        (dict(tone=SentimentLabel.ESCALATING, complaints={"Cost": "detected"}), "high"),  # 3+1
        (dict(tone=SentimentLabel.FRUSTRATED, escalation=EscalationLevel.HIGH), "high"),  # 2+2
        (dict(tone=None, escalation=EscalationLevel.CRITICAL), "medium"),  # 2
        (dict(tone=SentimentLabel.NEGATIVE, escalation=EscalationLevel.WATCH), "low"),  # 1
    ],
)
def test_churn_risk_follows_the_points_rule(call, risk):
    overall = _report(_call("a", 1, **call)).overall

    assert {"low": overall.churn_low, "medium": overall.churn_medium, "high": overall.churn_high} == {
        **{"low": 0, "medium": 0, "high": 0},
        risk: 1,
    }


def test_a_repeat_complaint_adds_to_churn_risk():
    report = _report(
        _call("first", 1, {"Cost": "resolved"}, customer="c:1", tone=SentimentLabel.NEGATIVE),
        _call("again", 5, {"Cost": "resolved"}, customer="c:1", tone=SentimentLabel.NEGATIVE),
    )

    # first: 1 (tone) = low. again: 1 (tone) + 2 (repeat) = medium.
    assert (report.overall.churn_low, report.overall.churn_medium) == (1, 1)


@pytest.mark.parametrize(
    ("call", "csat"),
    [
        (dict(tone=SentimentLabel.POSITIVE), 4.5),
        (dict(tone=SentimentLabel.POSITIVE, complaints={"Cost": "resolved"}), 5.0),
        (dict(tone=SentimentLabel.NEUTRAL, complaints={"Cost": "detected"}), 3.0),
        (dict(tone=SentimentLabel.NEGATIVE), 2.5),
        (dict(tone=SentimentLabel.FRUSTRATED, complaints={"Cost": "detected"}), 1.5),
        (
            dict(
                tone=SentimentLabel.ESCALATING,
                complaints={"Cost": "detected"},
                escalation=EscalationLevel.CRITICAL,
            ),
            1.0,  # 1.5 - 0.5 - 0.5 would be 0.5: never below 1
        ),
        (dict(tone=None, complaints={"Cost": "detected"}), None),
    ],
)
def test_the_csat_estimate_starts_from_the_tone_and_moves_with_the_outcome(call, csat):
    overall = _report(_call("a", 1, **call)).overall

    assert overall.csat == csat
    assert overall.rated_calls == (0 if csat is None else 1)


def test_csat_is_averaged_over_the_calls_that_have_a_tone():
    overall = _report(
        _call("a", 1, tone=SentimentLabel.POSITIVE),
        _call("b", 2, tone=SentimentLabel.NEGATIVE),
        _call("c", 3, tone=None),
    ).overall

    assert (overall.csat, overall.rated_calls, overall.negative_calls) == (3.5, 2, 1)


def test_the_scorecard_has_a_row_per_executive_busiest_first():
    report = _report(
        _call("a", 1, {"Cost": "detected"}, probed=("Cost",), executive="asha"),
        _call("b", 2, {"Cost": "detected"}, executive="ravi", escalation=EscalationLevel.HIGH),
        _call("c", 3, executive="ravi", tone=SentimentLabel.FRUSTRATED),
        _call("d", 4, executive=None),
    )

    assert [(e.executive_user_id, e.name, e.figures.calls) for e in report.executives] == [
        ("ravi", "Ravi", 2),
        ("asha", "Asha", 1),
        (None, "Not recorded", 1),
    ]
    ravi, asha, _ = (e.figures for e in report.executives)
    assert (asha.coverage_score, ravi.coverage_score) == (1.0, 0.0)
    assert (ravi.serious_escalations, ravi.negative_calls) == (1, 1)
    assert report.overall.calls == 4


def test_filters_narrow_the_calls_but_not_the_customers_history():
    report = _report(
        _call("ravi-first", 1, {"Cost": "detected"}, customer="c:1", executive="ravi"),
        _call("asha-again", 5, {"Cost": "detected"}, customer="c:1", executive="asha"),
        executive_user_id="asha",
    )

    # Only Asha's call is reported, and it is a repeat of the one Ravi took.
    assert (report.overall.calls, report.overall.repeat_complaints) == (1, 1)
    assert [e.name for e in report.executives] == ["Asha"]


def test_the_category_filter_keeps_the_calls_that_raised_it():
    report = _report(
        _call("a", 1, {"Cost": "detected"}),
        _call("b", 2, {"Hygiene": "detected"}),
        category="Hygiene",
    )

    assert report.overall.calls == 1


def test_too_long_or_too_large_a_report_is_refused():
    with pytest.raises(ReportError, match="at most 366 days"):
        _report(days=367)
    service = PerformanceService(_Calls(*[_call(str(n), n) for n in range(5)]), max_calls=3)
    with pytest.raises(ReportError, match="More than 3 calls"):
        service.report(ReportFilters(START, START + 30 * DAY))


# ---- The download ----

def test_the_scorecard_downloads_as_csv_and_excel():
    report = _report(
        _call("a", 1, {"Cost": "detected"}, probed=("Cost",), executive="asha", customer="c:1"),
        _call("b", 2, executive="=cmd", tone=None),
    )

    content, media_type, filename = export_scorecard(report, "csv", 330)
    rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))

    assert media_type.startswith("text/csv")
    assert filename.startswith("executive-scorecard-") and filename.endswith(".csv")
    assert rows[0][:5] == [
        "Executive", "Calls", "Complaints", "Complaints asked about", "Coverage score % (estimate)"
    ]
    # Equally busy executives are listed by name.
    # No tone and no complaints: empty cells, and a name is never run as a formula.
    assert rows[1][:9] == ["'=Cmd", "1", "0", "0", "", "", "0", "", ""]
    assert rows[2][:9] == ["Asha", "1", "1", "1", "100.0", "100.0", "1", "0.0", "3.0"]
    assert rows[3][0] == "All executives" and rows[3][1] == "2"

    workbook = load_workbook(io.BytesIO(export_scorecard(report, "xlsx")[0]))
    sheet = list(workbook["Scorecard"].iter_rows(values_only=True))
    assert sheet[2][:5] == ("Asha", 1, 1, 1, 100)
    assert len(sheet) == 4

    with pytest.raises(ValueError):
        export_scorecard(report, "pdf")


# ---- The API ----

class _NoComplaints(ComplaintDetectionProvider):
    def detect(self, conversation, learning_context=()):
        return []


class _Neutral(SentimentAnalysisProvider):
    def analyze(self, conversation, learning_context=()):
        return SentimentResult(SentimentLabel.NEUTRAL, 0.9, "Calm.")


class _NoQuestions(QuestionSuggestionProvider):
    def generate(self, context):
        return None


def test_the_api_serves_the_figures_to_supervisors_and_admins_only():
    services = build_api_services(
        _NoComplaints(), _Neutral(), _NoQuestions(), Settings(_env_file=None)  # type: ignore[call-arg]
    )
    client = TestClient(create_app(services))

    def sign_in(user_id: str, role: UserRole) -> dict[str, str]:
        user = User(user_id, f"{user_id}@example.com", "hash", role, True, time.time())
        services.user_repository.save(user)
        return {"Authorization": f"Bearer {create_access_token(user)}"}

    supervisor = sign_in("sup", UserRole.SUPERVISOR)
    # A call started in the app is taken by whoever is signed in.
    client.post("/api/v1/calls", json={"call_id": "c1", "start_time": None}, headers=supervisor)

    body = client.get("/api/v1/reports/performance", headers=supervisor).json()

    assert body["overall"]["calls"] == 1
    assert body["overall"]["coverage_score"] is None
    assert [(e["executive_user_id"], e["name"]) for e in body["executives"]] == [
        ("sup", "sup@example.com")
    ]
    assert body["executives"][0]["figures"]["churn_low"] == 1

    export = client.get(
        "/api/v1/reports/performance/export", params={"format": "csv"}, headers=supervisor
    )
    assert export.status_code == 200
    assert "sup@example.com" in export.content.decode("utf-8-sig")
    assert 'filename="executive-scorecard-' in export.headers["content-disposition"]

    assert client.get("/api/v1/reports/performance", headers=sign_in("adm", UserRole.ADMIN)).status_code == 200
    icr = sign_in("icr", UserRole.ICR)
    assert client.get("/api/v1/reports/performance", headers=icr).status_code == 403
    assert client.get("/api/v1/reports/performance/export", headers=icr).status_code == 403
    assert client.get("/api/v1/reports/performance").status_code == 401
    assert (
        client.get(
            "/api/v1/reports/performance/export", params={"format": "pdf"}, headers=supervisor
        ).status_code
        == 422
    )
