"""Complaint reports: what is counted, each filter, days and weeks, the
grouping of similar descriptions, who may see them, and the three exports.
The calls come from the in-memory stores and from SQL alike."""

import csv
import io
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.pool import StaticPool

from app.ai.complaint.provider import ComplaintDetectionProvider
from app.ai.question.provider import QuestionSuggestionProvider
from app.ai.sentiment.provider import (
    SentimentAnalysisProvider,
    SentimentLabel,
    SentimentResult,
)
from app.api.app_factory import create_app
from app.api.wiring import build_api_services
from app.core.config import Settings
from app.domain.complaint_coverage import ComplaintCoverageStatus
from app.domain.complaint_lifecycle import ComplaintLifecycleRecord, ComplaintLifecycleStatus
from app.domain.complaint_lifecycle_repository import InMemoryComplaintLifecycleRepository
from app.infrastructure.database.repositories.complaint_lifecycle_repository import (
    PostgresComplaintLifecycleRepository,
)
from app.domain.conversation import CallDirection, Conversation
from app.domain.conversation_coverage import ConversationCoverage
from app.domain.location import InMemoryLocationRepository, Location
from app.domain.post_call_summary import ComplaintSummary, PostCallSummary
from app.domain.user import User, UserRole
from app.domain.user_repository import InMemoryUserRepository
from app.infrastructure.database.base import Base
from app.infrastructure.database.engine import build_session_factory
from app.infrastructure.database.repositories.conversation_coverage_repository import (
    PostgresConversationCoverageRepository,
)
from app.infrastructure.database.repositories.conversation_repository import (
    PostgresConversationRepository,
)
from app.infrastructure.database.repositories.location_repository import (
    PostgresLocationRepository,
)
from app.infrastructure.database.repositories.post_call_summary_repository import (
    PostgresPostCallSummaryRepository,
)
from app.infrastructure.database.repositories.report_source import PostgresReportSource
from app.infrastructure.database.repositories.user_repository import PostgresUserRepository
from app.security.jwt import create_access_token
from app.services.in_memory_conversation_coverage_repository import (
    InMemoryConversationCoverageRepository,
)
from app.services.in_memory_conversation_repository import InMemoryConversationRepository
from app.services.post_call_summary_repository import InMemoryPostCallSummaryRepository
from app.services.report_export import export_report
from app.services.reporting import (
    InMemoryReportSource,
    ReportError,
    ReportFilters,
    ReportService,
)

DAY = 86400.0
IST = 330  # minutes ahead of UTC
# Monday 5 January 2026, 00:00 in India.
MONDAY = 1767551400.0
_STEPS = ("detect", "probe", "cover")


def _coverage(call_id: str, complaints: dict[str, str]) -> ConversationCoverage:
    """Complaints as {category: status at the end of the call}."""
    coverage = ConversationCoverage(call_id)
    for category, status in complaints.items():
        complaint = coverage.add(category)
        if status == "not_raised":
            continue
        for step in _STEPS[: {"detected": 1, "probed": 2}.get(status, 3)]:
            getattr(complaint, step)()
        if status == "resolved":
            complaint.resolve()
        elif status == "unresolved":
            complaint.mark_unresolved()
    return coverage


def _summary(call_id: str, sentiment: SentimentLabel, described: dict[str, str]) -> PostCallSummary:
    return PostCallSummary(
        call_id=call_id,
        overall_summary="A call.",
        languages=("en",),
        sentiment=SentimentResult(sentiment, 0.9, "What was said."),
        complaints=tuple(
            ComplaintSummary(category, text, ComplaintCoverageStatus.DETECTED, "Evidence.")
            for category, text in described.items()
        ),
        unresolved_issues=(),
        actions_promised=(),
        follow_up_required=False,
        customer_summary="A summary.",
    )


class _Stores:
    def __init__(self, session_factory=None) -> None:
        if session_factory is None:
            self.conversations = InMemoryConversationRepository()
            self.coverages = InMemoryConversationCoverageRepository()
            self.summaries = InMemoryPostCallSummaryRepository()
            self.locations = InMemoryLocationRepository()
            self.users = InMemoryUserRepository()
            self.tracked = InMemoryComplaintLifecycleRepository()
            self.source = InMemoryReportSource(
                self.conversations,
                self.coverages,
                self.summaries,
                self.locations,
                self.users,
                self.tracked,
            )
        else:
            self.conversations = PostgresConversationRepository(session_factory)
            self.coverages = PostgresConversationCoverageRepository(session_factory)
            self.summaries = PostgresPostCallSummaryRepository(session_factory)
            self.locations = PostgresLocationRepository(session_factory)
            self.users = PostgresUserRepository(session_factory)
            self.tracked = PostgresComplaintLifecycleRepository(session_factory)
            self.source = PostgresReportSource(session_factory)
        self.service = ReportService(self.source, self.locations, self.users)

    def call(
        self,
        call_id: str,
        day: float,
        complaints: dict[str, str] | None = None,
        *,
        location: str | None = None,
        executive: str | None = None,
        direction: CallDirection | None = CallDirection.INBOUND,
        sentiment: SentimentLabel | None = SentimentLabel.NEUTRAL,
        described: dict[str, str] | None = None,
    ) -> None:
        """A call `day` days after MONDAY; sentiment None: no summary yet."""
        self.conversations.add(
            Conversation(
                call_id,
                start_time=MONDAY + day * DAY,
                direction=direction,
                location_id=location,
                executive_user_id=executive,
            )
        )
        if complaints is not None:
            self.coverages.save(_coverage(call_id, complaints))
        if sentiment is not None:
            self.summaries.add_if_absent(_summary(call_id, sentiment, described or {}))

    def report(self, days: float = 14, bucket=None, **filters):
        return self.service.complaint_report(
            ReportFilters(MONDAY, MONDAY + days * DAY, **filters), bucket, IST
        )


@pytest.fixture(params=["memory", "sql"])
def stores(request) -> _Stores:
    if request.param == "memory":
        s = _Stores()
    else:
        engine = create_engine(
            "sqlite:///:memory:",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(engine)
        request.addfinalizer(engine.dispose)
        s = _Stores(build_session_factory(engine))
    s.locations.save(Location("chennai", "Chennai", "+914440000001", True, 1.0))
    s.locations.save(Location("madurai", "Madurai", "+914520000002", True, 1.0))
    s.users.save(
        User("asha", "asha@example.com", "hash", UserRole.ICR, True, 1.0, display_name="Asha")
    )
    s.users.save(User("ravi", "ravi@example.com", "hash", UserRole.ICR, True, 1.0))

    s.call(
        "c1",
        0.5,
        {"Cost": "resolved", "Hygiene": "detected"},
        location="chennai",
        executive="asha",
        sentiment=SentimentLabel.NEGATIVE,
        described={
            "Cost": "Customer was charged more than the estimate.",
            "Hygiene": "The car was returned dirty.",
        },
    )
    s.call(
        "c2",
        1.5,
        {"Cost": "unresolved"},
        location="chennai",
        executive="ravi",
        sentiment=SentimentLabel.NEGATIVE,
        described={"Cost": "The customer says he was charged more than the estimated amount."},
    )
    s.call(
        "c3",
        8.2,
        {"Cost": "probed", "Hospitality": "not_raised"},
        location="madurai",
        executive="asha",
        direction=CallDirection.OUTBOUND,
        described={"Cost": "Labour charges were not explained."},
    )
    # Still without a summary: counted, not yet described.
    s.call("c4", 9.0, {"Hygiene": "detected"}, sentiment=None)
    # No complaints at all.
    s.call("c5", 2.0, location="madurai", executive="ravi", sentiment=SentimentLabel.POSITIVE)
    # Outside the fortnight.
    s.call("before", -1.0, {"Cost": "detected"}, location="chennai")
    s.call("after", 14.0, {"Cost": "detected"}, location="chennai")
    return s


# ---- What is counted ----

def test_totals_count_calls_and_the_complaints_raised_on_them(stores):
    report = stores.report()

    assert (report.total_calls, report.calls_with_complaints, report.total_complaints) == (5, 4, 5)
    assert [(c.category, c.complaints, c.resolved) for c in report.categories] == [
        ("Cost", 3, 1),
        ("Hygiene", 2, 0),
    ]
    # A category that never came up on the call is not a complaint.
    assert "Hospitality" not in {row.category for row in report.rows}


def test_the_range_includes_its_start_and_excludes_its_end(stores):
    assert stores.report(days=14).total_calls == 5
    assert stores.report(days=14.0001).total_calls == 6
    assert {r.call_id for r in stores.report(days=1).rows} == {"c1"}


def test_a_tracked_complaint_counts_with_the_status_it_has_now(stores):
    # c2's cost complaint was left unresolved on the call and closed later;
    # c1's hygiene complaint is tracked and still open.
    for call_id, category, status in (
        ("c2", "Cost", ComplaintLifecycleStatus.RESOLVED),
        ("c1", "Hygiene", ComplaintLifecycleStatus.FOLLOW_UP),
    ):
        stores.tracked.save(
            ComplaintLifecycleRecord(f"{call_id}:{category}", call_id, category, status, 1.0, 2.0)
        )

    report = stores.report()

    assert [(c.category, c.resolved) for c in report.categories] == [("Cost", 2), ("Hygiene", 0)]
    statuses = {(r.call_id, r.category): r.status for r in report.rows}
    assert statuses["c2", "Cost"] == "resolved"
    assert statuses["c1", "Hygiene"] == "follow_up"
    # Not tracked: as the call left it.
    assert statuses["c3", "Cost"] == "probed"


def test_a_description_that_says_nothing_is_not_a_root_cause(stores):
    stores.call(
        "plain",
        6.0,
        {"TAT": "detected"},
        described={"TAT": "TAT complaint identified during the call (status: detected)."},
    )

    tat = next(c for c in stores.report().root_causes if c.category == "TAT")

    assert (tat.complaints, tat.undescribed, tat.themes) == (1, 1, ())


# ---- Filters ----

@pytest.mark.parametrize(
    ("filters", "calls", "complaints"),
    [
        ({"location_id": "chennai"}, 2, 3),
        ({"location_id": "nowhere"}, 0, 0),
        ({"executive_user_id": "asha"}, 2, 3),
        ({"direction": CallDirection.OUTBOUND}, 1, 1),
        ({"sentiment": SentimentLabel.NEGATIVE}, 2, 3),
        ({"sentiment": SentimentLabel.POSITIVE}, 1, 0),
        ({"location_id": "chennai", "executive_user_id": "ravi"}, 1, 1),
    ],
)
def test_each_filter_narrows_the_calls(stores, filters, calls, complaints):
    report = stores.report(**filters)

    assert (report.total_calls, report.total_complaints) == (calls, complaints)


def test_the_category_filter_keeps_only_that_categorys_complaints(stores):
    report = stores.report(category="Hygiene")

    # c1 also has a Cost complaint: the call counts, that complaint does not.
    assert (report.total_calls, report.total_complaints) == (2, 2)
    assert [c.category for c in report.categories] == ["Hygiene"]
    assert {row.call_id for row in report.rows} == {"c1", "c4"}


def test_the_filters_ids_are_named_for_headings(stores):
    report = stores.report(location_id="chennai", executive_user_id="ravi")

    assert (report.location_name, report.executive_name) == ("Chennai", "ravi@example.com")
    assert stores.report().location_name is None


# ---- Trend ----

def test_a_short_range_is_counted_by_day(stores):
    report = stores.report()

    assert report.bucket == "day"
    assert len(report.bucket_starts) == 14
    assert report.bucket_starts[0] == MONDAY
    assert report.bucket_starts[1] - report.bucket_starts[0] == DAY
    cost, hygiene = report.trend
    assert cost.counts == (1, 1, 0, 0, 0, 0, 0, 0, 1, 0, 0, 0, 0, 0)
    assert hygiene.counts == (1, 0, 0, 0, 0, 0, 0, 0, 0, 1, 0, 0, 0, 0)


def test_weeks_run_monday_to_sunday_by_the_readers_clock(stores):
    report = stores.report(bucket="week")

    assert report.bucket_starts == (MONDAY, MONDAY + 7 * DAY)
    assert [series.counts for series in report.trend] == [(2, 1), (1, 1)]


def test_a_day_starts_at_the_readers_midnight(stores):
    # 02:00 on Monday in India is still Sunday evening by UTC.
    stores.call("early", 2 / 24, {"TAT": "detected"})
    in_india = stores.report(category="TAT")
    in_utc = stores.service.complaint_report(
        ReportFilters(MONDAY, MONDAY + 14 * DAY, category="TAT"), None, 0
    )

    assert in_india.bucket_starts[0] == MONDAY
    assert in_india.trend[0].counts[0] == 1
    # By UTC the same fortnight starts on Sunday and touches fifteen days.
    assert in_utc.bucket_starts[0] == MONDAY - 18.5 * 3600
    assert len(in_utc.bucket_starts) == 15
    assert in_utc.trend[0].counts[0] == 1


def test_a_long_range_is_counted_by_week_unless_asked_otherwise(stores):
    assert stores.report(days=90).bucket == "week"
    assert stores.report(days=90, bucket="day").bucket == "day"
    assert len(stores.report(days=90, bucket="day").bucket_starts) == 90


# ---- Heatmap ----

def test_the_heatmap_counts_each_category_per_location(stores):
    report = stores.report()

    assert [(l.location_id, l.name) for l in report.locations] == [
        ("chennai", "Chennai"),
        ("madurai", "Madurai"),
        (None, "No location"),
    ]
    assert [(row.category, row.counts) for row in report.heatmap] == [
        ("Cost", (2, 1, 0)),
        ("Hygiene", (1, 0, 1)),
    ]


# ---- Root causes ----

def test_complaints_in_similar_words_are_grouped(stores):
    cost, hygiene = stores.report().root_causes

    assert (cost.category, cost.complaints, cost.resolved, cost.undescribed) == ("Cost", 3, 1, 0)
    assert [(t.text, t.complaints, t.call_ids) for t in cost.themes] == [
        ("Customer was charged more than the estimate.", 2, ("c2", "c1")),
        ("Labour charges were not explained.", 1, ("c3",)),
    ]
    # c4 has no summary yet, so nothing describes its complaint.
    assert (hygiene.complaints, hygiene.undescribed) == (2, 1)
    assert [t.text for t in hygiene.themes] == ["The car was returned dirty."]


def test_only_the_largest_themes_and_a_few_calls_each_are_listed(stores):
    for index in range(8):
        stores.call(
            f"wait-{index}",
            3.0 + index / 100,
            {"TAT": "detected"},
            described={"TAT": "The vehicle was delivered two days late."},
        )
    for index, text in enumerate(
        ["Nobody called back.", "Pickup driver arrived late.", "Wrong invoice sent.",
         "Parts ordered twice.", "Appointment slot lost.", "Washing skipped entirely."]
    ):
        stores.call(f"one-{index}", 4.0 + index / 100, {"TAT": "detected"}, described={"TAT": text})

    tat = next(c for c in stores.report().root_causes if c.category == "TAT")

    assert tat.complaints == 14
    assert len(tat.themes) == 5
    assert tat.themes[0].complaints == 8
    # The most recent calls of the theme.
    assert tat.themes[0].call_ids == ("wait-7", "wait-6", "wait-5", "wait-4", "wait-3")


# ---- Limits ----

def test_a_report_that_is_too_large_or_too_long_is_refused(stores):
    small = ReportService(stores.source, max_calls=4)
    with pytest.raises(ReportError, match="More than 4 calls"):
        small.complaint_report(ReportFilters(MONDAY, MONDAY + 14 * DAY))
    assert small.complaint_report(ReportFilters(MONDAY, MONDAY + 2 * DAY)).total_calls == 2

    with pytest.raises(ReportError, match="at most 366 days"):
        ReportFilters(MONDAY, MONDAY + 367 * DAY)
    with pytest.raises(ReportError, match="after its start"):
        ReportFilters(MONDAY, MONDAY)


def test_an_empty_report_still_has_its_days(stores):
    report = stores.report(location_id="nowhere")

    assert report.categories == report.trend == report.heatmap == report.root_causes == ()
    assert report.locations == ()
    assert len(report.bucket_starts) == 14


# ---- Exports ----

def test_the_csv_lists_one_complaint_per_row(stores):
    content, media_type, filename = export_report(stores.report(), "csv")

    assert media_type.startswith("text/csv")
    assert filename == "complaint-report-2026-01-05.csv"
    rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    assert rows[0][:4] == ["Call ID", "Call date", "Location", "Executive"]
    assert len(rows) == 1 + 5
    assert rows[1] == [
        "c1",
        "2026-01-05 12:00",
        "Chennai",
        "Asha",
        "Incoming",
        "Negative",
        "Cost",
        "resolved",
        "Customer was charged more than the estimate.",
    ]
    # No summary, no location, no executive: empty cells.
    assert rows[5][:9] == ["c4", "2026-01-14 00:00", "", "", "Incoming", "", "Hygiene", "detected", ""]


def test_the_excel_file_has_a_sheet_per_section(stores):
    content, media_type, filename = export_report(
        stores.report(location_id="chennai"), "xlsx"
    )

    assert filename.endswith(".xlsx")
    workbook = load_workbook(io.BytesIO(content))
    assert workbook.sheetnames == ["Summary", "Trend", "By location", "Root causes", "Complaints"]
    summary = {row[0]: row[1] for row in workbook["Summary"].iter_rows(values_only=True)}
    assert summary["Location"] == "Chennai"
    assert summary["Calls from"] == "05 Jan 2026"
    assert summary["Calls to"] == "18 Jan 2026"
    assert (summary["Calls"], summary["Complaints"], summary["Cost"]) == (2, 3, 2)
    trend = list(workbook["Trend"].iter_rows(values_only=True))
    assert trend[0][:3] == ("Category", "05 Jan 2026", "06 Jan 2026")
    assert trend[1][:3] == ("Cost", 1, 1)
    assert list(workbook["By location"].iter_rows(values_only=True)) == [
        ("Category", "Chennai"),
        ("Cost", 2),
        ("Hygiene", 1),
    ]
    causes = list(workbook["Root causes"].iter_rows(values_only=True))
    assert causes[1] == ("Cost", "Customer was charged more than the estimate.", 2, "c2, c1")
    assert workbook["Complaints"].max_row == 1 + 3


def test_text_that_looks_like_a_formula_is_not_run_by_a_spreadsheet(stores):
    stores.call("x", 5.0, {"Cost": "detected"}, described={"Cost": "=HYPERLINK(\"http://x\")"})

    csv_rows = list(
        csv.reader(io.StringIO(export_report(stores.report(), "csv")[0].decode("utf-8-sig")))
    )
    workbook = load_workbook(io.BytesIO(export_report(stores.report(), "xlsx")[0]))

    assert "'=HYPERLINK(\"http://x\")" in [row[8] for row in csv_rows]
    cells = [row[8] for row in workbook["Complaints"].iter_rows(values_only=True)]
    assert "'=HYPERLINK(\"http://x\")" in cells


def test_the_pdf_is_a_pdf_with_the_reports_sections(stores):
    from reportlab.lib.utils import ImageReader  # noqa: F401  (the library is installed)

    content, media_type, filename = export_report(stores.report(bucket="week"), "pdf")
    empty, _, _ = export_report(stores.report(location_id="nowhere"), "pdf")

    assert media_type == "application/pdf"
    assert filename == "complaint-report-2026-01-05.pdf"
    assert content.startswith(b"%PDF-") and content.rstrip().endswith(b"%%EOF")
    assert empty.startswith(b"%PDF-")
    assert len(content) > len(empty)


def test_a_long_report_fits_the_pdf_page(stores):
    for index in range(12):
        stores.locations.save(Location(f"l{index}", f"Branch {index:02d}", f"+9144000001{index:02d}", True, 1.0))
        stores.call(f"b{index}", 20.0 + index, {"Cost": "detected"}, location=f"l{index}")

    content, _, _ = export_report(stores.report(days=60), "pdf")

    assert content.startswith(b"%PDF-")


def test_an_unknown_export_format_is_refused(stores):
    with pytest.raises(ValueError, match="Unsupported export format"):
        export_report(stores.report(), "docx")


# ---- The API ----

class _NoComplaints(ComplaintDetectionProvider):
    def detect(self, conversation):
        return []


class _Neutral(SentimentAnalysisProvider):
    def analyze(self, conversation):
        return SentimentResult(label=SentimentLabel.NEUTRAL, confidence=0.9, evidence="Calm.")


class _NoQuestions(QuestionSuggestionProvider):
    def generate(self, context):
        return None


def _api():
    stores = _Stores()
    services = build_api_services(
        _NoComplaints(),
        _Neutral(),
        _NoQuestions(),
        Settings(_env_file=None),  # type: ignore[call-arg]
        conversation_repository=stores.conversations,
        coverage_repository=stores.coverages,
        post_call_summary_repository=stores.summaries,
        location_repository=stores.locations,
        user_repository=stores.users,
    )
    client = TestClient(create_app(services))

    def sign_in(user_id: str, role: UserRole) -> dict[str, str]:
        user = User(user_id, f"{user_id}@example.com", "hash", role, True, time.time())
        stores.users.save(user)
        return {"Authorization": f"Bearer {create_access_token(user)}"}

    return client, stores, sign_in


RANGE = {"started_from": MONDAY, "started_to": MONDAY + 14 * DAY, "tz_offset_minutes": IST}


def test_supervisors_and_admins_see_reports_and_executives_do_not():
    client, stores, sign_in = _api()
    stores.locations.save(Location("chennai", "Chennai", "+914440000001", True, 1.0))
    stores.call("c1", 0.5, {"Cost": "resolved"}, location="chennai",
                described={"Cost": "Charged more than the estimate."})

    for role in (UserRole.SUPERVISOR, UserRole.ADMIN):
        response = client.get(
            "/api/v1/reports/complaints", params=RANGE, headers=sign_in(role.value, role)
        )
        assert response.status_code == 200
        body = response.json()
        assert (body["total_calls"], body["total_complaints"], body["bucket"]) == (1, 1, "day")
        assert body["categories"] == [{"category": "Cost", "complaints": 1, "resolved": 1}]
        assert body["locations"] == [{"location_id": "chennai", "name": "Chennai"}]
        assert body["heatmap"] == [{"category": "Cost", "counts": [1]}]
        assert body["root_causes"][0]["themes"][0]["call_ids"] == ["c1"]
        assert len(body["bucket_starts"]) == len(body["trend"][0]["counts"]) == 14
        # The complaint rows are for the exports only.
        assert "rows" not in body

    icr = sign_in("icr", UserRole.ICR)
    assert client.get("/api/v1/reports/complaints", params=RANGE, headers=icr).status_code == 403
    assert (
        client.get("/api/v1/reports/complaints/export", params=RANGE, headers=icr).status_code
        == 403
    )
    assert client.get("/api/v1/reports/complaints", params=RANGE).status_code == 401


def test_the_api_applies_filters_and_refuses_bad_ones():
    client, stores, sign_in = _api()
    headers = sign_in("sup", UserRole.SUPERVISOR)
    stores.call("c1", 0.5, {"Cost": "detected"}, sentiment=SentimentLabel.NEGATIVE)
    stores.call("c2", 1.5, {"Hygiene": "detected"}, direction=CallDirection.OUTBOUND)

    def get(**params):
        return client.get(
            "/api/v1/reports/complaints", params={**RANGE, **params}, headers=headers
        )

    assert get(sentiment="NEGATIVE").json()["total_calls"] == 1
    assert get(direction="outbound").json()["categories"][0]["category"] == "Hygiene"
    assert get(category="Cost", bucket="week").json()["bucket"] == "week"
    assert get(category="  ").json()["total_calls"] == 2

    too_long = get(started_to=MONDAY + 400 * DAY)
    assert too_long.status_code == 422
    assert "at most 366 days" in too_long.json()["detail"]
    assert get(started_to=MONDAY - 1).status_code == 422
    assert get(sentiment="ANGRY").status_code == 422
    assert get(bucket="month").status_code == 422
    assert get(tz_offset_minutes=5000).status_code == 422


def test_without_dates_the_report_covers_the_last_thirty_days():
    client, stores, sign_in = _api()
    headers = sign_in("sup", UserRole.SUPERVISOR)
    now = time.time()
    stores.conversations.add(Conversation("recent", start_time=now - 5 * DAY))
    stores.conversations.add(Conversation("old", start_time=now - 45 * DAY))

    body = client.get("/api/v1/reports/complaints", headers=headers).json()

    assert body["total_calls"] == 1
    assert body["filters"]["started_to"] - body["filters"]["started_from"] == pytest.approx(
        30 * DAY
    )


@pytest.mark.parametrize(
    ("file_format", "media_type", "starts"),
    [
        ("csv", "text/csv", b"\xef\xbb\xbfCall ID"),
        ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", b"PK"),
        ("pdf", "application/pdf", b"%PDF-"),
    ],
)
def test_each_export_downloads_what_the_filters_show(file_format, media_type, starts):
    client, stores, sign_in = _api()
    headers = sign_in("admin", UserRole.ADMIN)
    stores.call("c1", 0.5, {"Cost": "detected"}, described={"Cost": "Charged too much."})
    stores.call("c2", 1.5, {"Hygiene": "detected"})

    response = client.get(
        "/api/v1/reports/complaints/export",
        params={**RANGE, "format": file_format, "category": "Cost"},
        headers=headers,
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(media_type)
    assert (
        response.headers["content-disposition"]
        == f'attachment; filename="complaint-report-2026-01-05.{file_format}"'
    )
    assert response.content.startswith(starts)
    if file_format == "csv":
        text = response.content.decode("utf-8-sig")
        assert "Charged too much." in text and "Hygiene" not in text

    bad = client.get(
        "/api/v1/reports/complaints/export", params={**RANGE, "format": "docx"}, headers=headers
    )
    assert bad.status_code == 422
