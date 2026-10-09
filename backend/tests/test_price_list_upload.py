"""Reading an uploaded price-list workbook stays fast and bounded, and the
price list keeps working while the database is down."""

import io
import time

import openpyxl
import pytest

from app.estimation.default_pricing import DEFAULT_PRICE_LIST_ROWS
from app.estimation import price_list_io
from app.estimation.price_list_io import (
    EXPORT_HEADERS,
    MAX_COLUMNS,
    PriceListFileError,
    export_xlsx,
    parse_price_list,
)
from app.services.price_list_service import PriceListService


def workbook_bytes(rows: int, stray_column: int | None = None) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Price list"
    sheet.append(list(EXPORT_HEADERS))
    for i in range(rows):
        sheet.append([None, f"Service {i % 20}", f"Part {i}", 1, 100 + i, 18, 1, 1, None, None])
    if stray_column is not None:
        sheet.cell(row=2, column=stray_column, value="stray")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_a_stray_cell_in_the_last_excel_column_does_not_slow_reading():
    content = workbook_bytes(2000, stray_column=16384)  # column XFD

    started = time.perf_counter()
    parsed = parse_price_list("prices.xlsx", content)

    assert time.perf_counter() - started < 10
    assert len(parsed.rows) == 2000
    assert parsed.rows[0].service == "Service 0"


def test_columns_after_the_first_ones_are_not_read():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Part"] + [None] * MAX_COLUMNS + ["Service"])
    sheet.append(["Oil filter"] + [None] * MAX_COLUMNS + ["Oil Change"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    with pytest.raises(PriceListFileError, match="'Service' column"):
        parse_price_list("prices.xlsx", buffer.getvalue())


def test_a_workbook_that_unpacks_too_large_is_refused(monkeypatch):
    monkeypatch.setattr(price_list_io, "MAX_UNPACKED_XLSX_BYTES", 1000)

    with pytest.raises(PriceListFileError, match="too large once opened"):
        parse_price_list("prices.xlsx", workbook_bytes(50))


def test_a_file_that_is_not_a_zip_is_reported_as_unreadable():
    with pytest.raises(PriceListFileError, match="could not be opened"):
        parse_price_list("prices.xlsx", b"not a workbook")


def test_the_exported_list_reads_back_unchanged():
    parsed = parse_price_list("prices.xlsx", export_xlsx(DEFAULT_PRICE_LIST_ROWS))

    assert parsed.errors == ()
    assert parsed.rows == DEFAULT_PRICE_LIST_ROWS


class FlakyRepository:
    """A price-list repository whose database can be switched off."""

    def __init__(self) -> None:
        self.down = False
        self.reads = 0

    def get_active(self):
        self.reads += 1
        if self.down:
            raise RuntimeError("database down")
        return None  # nothing saved: the sample price list


def test_while_the_database_is_down_it_is_retried_once_per_cache_time():
    repository = FlakyRepository()
    now = [0.0]
    service = PriceListService(repository, cache_seconds=30.0, clock=lambda: now[0])
    first = service.pricing()

    repository.down = True
    now[0] = 31.0
    for _ in range(5):
        assert service.pricing() is first  # the last list read
    assert repository.reads == 2  # one failed retry, not five

    now[0] = 62.0
    repository.down = False
    service.pricing()
    assert repository.reads == 3
