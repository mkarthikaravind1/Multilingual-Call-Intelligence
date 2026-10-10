"""A complaint report as a file: CSV (the complaints, one per row), Excel
(the same plus a sheet per section) or PDF (the sections, to print)."""

import csv
import io
from datetime import datetime, timedelta, timezone
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.services.reporting import Report

EXPORT_FORMATS = ("csv", "xlsx", "pdf")
_MEDIA_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}
_COMPLAINT_COLUMNS = (
    "Call ID",
    "Call date",
    "Location",
    "Executive",
    "Direction",
    "Tone",
    "Category",
    "Status",
    "Description",
)
_DIRECTIONS = {"inbound": "Incoming", "outbound": "Outgoing"}
# The PDF's tables keep to what fits a page width.
_PDF_MAX_TREND_COLUMNS = 16
_PDF_MAX_LOCATION_COLUMNS = 10
# Spreadsheets run a cell that starts with one of these as a formula.
_FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")


def export_report(report: Report, file_format: str) -> tuple[bytes, str, str]:
    """(content, media type, filename)."""
    if file_format not in EXPORT_FORMATS:
        raise ValueError(f"Unsupported export format: {file_format!r}.")
    content = {"csv": _csv, "xlsx": _xlsx, "pdf": _pdf}[file_format](report)
    day = _local(report, report.filters.started_from).strftime("%Y-%m-%d")
    return content, _MEDIA_TYPES[file_format], f"complaint-report-{day}.{file_format}"


# ---- Shared ----


def _local(report: Report, at: float) -> datetime:
    """The moment by the reader's clock."""
    zone = timezone(timedelta(minutes=report.tz_offset_minutes))
    return datetime.fromtimestamp(at, zone)


def _bucket_labels(report: Report) -> list[str]:
    prefix = "Week of " if report.bucket == "week" else ""
    return [f"{prefix}{_local(report, at):%d %b %Y}" for at in report.bucket_starts]


def _filter_lines(report: Report) -> list[tuple[str, str]]:
    filters = report.filters
    # The range's end is exclusive: its last day is the one before.
    last_day = _local(report, filters.started_to - 1)
    lines = [
        ("Calls from", f"{_local(report, filters.started_from):%d %b %Y}"),
        ("Calls to", f"{last_day:%d %b %Y}"),
        ("Location", report.location_name or ("All" if filters.location_id is None else "?")),
        (
            "Executive",
            report.executive_name or ("All" if filters.executive_user_id is None else "?"),
        ),
        ("Category", filters.category or "All"),
        ("Tone", "All" if filters.sentiment is None else filters.sentiment.value.capitalize()),
        (
            "Direction",
            "All" if filters.direction is None else _DIRECTIONS[filters.direction.value],
        ),
    ]
    return lines


def _total_lines(report: Report) -> list[tuple[str, int]]:
    return [
        ("Calls", report.total_calls),
        ("Calls with a complaint", report.calls_with_complaints),
        ("Complaints", report.total_complaints),
    ]


def _complaint_rows(report: Report) -> list[list[str]]:
    return [
        [
            row.call_id,
            f"{_local(report, row.start_time):%Y-%m-%d %H:%M}",
            row.location_name or "",
            row.executive_name or "",
            "" if row.direction is None else _DIRECTIONS[row.direction.value],
            "" if row.sentiment is None else row.sentiment.value.capitalize(),
            row.category,
            row.status.replace("_", " "),
            row.description or "",
        ]
        for row in report.rows
    ]


def _safe(value):
    """Text that a spreadsheet will show rather than run."""
    if isinstance(value, str) and value.startswith(_FORMULA_STARTS):
        return "'" + value
    return value


# ---- CSV ----


def _csv(report: Report) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(_COMPLAINT_COLUMNS)
    for row in _complaint_rows(report):
        writer.writerow([_safe(cell) for cell in row])
    # With a byte-order mark, so Excel reads it as UTF-8.
    return buffer.getvalue().encode("utf-8-sig")


# ---- Excel ----


def _sheet(workbook: Workbook, title: str, header: list[str], rows: list[list]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([_safe(cell) for cell in row])
    sheet.freeze_panes = "A2"
    for index, name in enumerate(header, start=1):
        longest = max([len(str(name))] + [len(str(row[index - 1])) for row in rows[:200]])
        sheet.column_dimensions[get_column_letter(index)].width = min(60, max(10, longest + 2))


def _xlsx(report: Report) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)

    summary = [[name, value] for name, value in _filter_lines(report)]
    summary += [["", ""]] + [[name, value] for name, value in _total_lines(report)]
    summary += [["", ""], ["Category", "Complaints", "Resolved"]]
    summary += [[c.category, c.complaints, c.resolved] for c in report.categories]
    _sheet(workbook, "Summary", ["Complaint report", ""], summary)

    _sheet(
        workbook,
        "Trend",
        ["Category"] + _bucket_labels(report),
        [[series.category, *series.counts] for series in report.trend],
    )
    _sheet(
        workbook,
        "By location",
        ["Category"] + [location.name for location in report.locations],
        [[row.category, *row.counts] for row in report.heatmap],
    )
    _sheet(
        workbook,
        "Root causes",
        ["Category", "What the complaints say (similar wording grouped)", "Complaints", "Calls"],
        [
            [cause.category, theme.text, theme.complaints, ", ".join(theme.call_ids)]
            for cause in report.root_causes
            for theme in cause.themes
        ],
    )
    _sheet(workbook, "Complaints", list(_COMPLAINT_COLUMNS), _complaint_rows(report))

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---- PDF ----


def _pdf(report: Report) -> bytes:
    # Imported here: only this export needs the PDF library.
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    body = styles["BodyText"]
    small = styles["BodyText"].clone("small", fontSize=8, leading=10)

    def text(value, style=small) -> Paragraph:
        return Paragraph(escape(str(value)), style)

    def table(rows: list[list], widths: list[float] | None = None) -> Table:
        made = Table([[text(cell) for cell in row] for row in rows], colWidths=widths, repeatRows=1)
        made.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e9f0f7")),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c5d0dc")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        return made

    story: list = [Paragraph("Complaint report", styles["Title"])]
    story.append(
        text(
            " · ".join(f"{name}: {value}" for name, value in _filter_lines(report)),
            body,
        )
    )
    story.append(
        text(" · ".join(f"{name}: {value}" for name, value in _total_lines(report)), body)
    )

    def section(title: str) -> None:
        story.append(Spacer(1, 5 * mm))
        story.append(Paragraph(escape(title), styles["Heading2"]))

    if not report.categories:
        story.append(Spacer(1, 5 * mm))
        story.append(text("No complaints match these filters.", body))

    if report.categories:
        section("Complaints by category")
        story.append(
            table(
                [["Category", "Complaints", "Resolved"]]
                + [[c.category, c.complaints, c.resolved] for c in report.categories],
                [90 * mm, 35 * mm, 45 * mm],
            )
        )

        # The most recent periods, when there are more than fit the page.
        labels = _bucket_labels(report)
        shown = slice(max(0, len(labels) - _PDF_MAX_TREND_COLUMNS), len(labels))
        section(
            "Trend by "
            + report.bucket
            + (
                f" (last {_PDF_MAX_TREND_COLUMNS} of {len(labels)}; all are in the Excel export)"
                if len(labels) > _PDF_MAX_TREND_COLUMNS
                else ""
            )
        )
        story.append(
            table(
                [["Category"] + [label.replace("Week of ", "") for label in labels[shown]]]
                + [[series.category, *series.counts[shown]] for series in report.trend]
            )
        )

        locations = report.locations[:_PDF_MAX_LOCATION_COLUMNS]
        section(
            "Category by location"
            + (
                f" (first {len(locations)} of {len(report.locations)} locations)"
                if len(report.locations) > len(locations)
                else ""
            )
        )
        story.append(
            table(
                [["Category"] + [location.name for location in locations]]
                + [[row.category, *row.counts[: len(locations)]] for row in report.heatmap]
            )
        )

        section("Root causes (complaints in similar words, grouped by a fixed rule)")
        for cause in report.root_causes:
            story.append(
                text(
                    f"{cause.category}: {cause.complaints} complaints, "
                    f"{cause.resolved} resolved"
                    + (f", {cause.undescribed} not described in a call summary" if cause.undescribed else ""),
                    body,
                )
            )
            if cause.themes:
                story.append(
                    table(
                        [["What the complaints say", "Complaints"]]
                        + [[theme.text, theme.complaints] for theme in cause.themes],
                        [215 * mm, 30 * mm],
                    )
                )
            story.append(Spacer(1, 2 * mm))

    buffer = io.BytesIO()
    SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title="Complaint report",
    ).build(story)
    return buffer.getvalue()
