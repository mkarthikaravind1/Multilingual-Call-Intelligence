"""Reading the supervisor's price-list spreadsheet (.xlsx or .csv) into
rows, and writing rows back out in the same layout.

One row per vehicle model + service + part. Columns are matched by name
(case and punctuation do not matter, a few common alternatives are
accepted), so their order is free and extra columns are ignored.
"""

import csv
import io
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from typing import Any

from app.estimation.price_list import MAX_ROWS, PriceListIssue, PriceListRow

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
_HEADER_SEARCH_ROWS = 10

# (field, header written on export, other accepted headers)
_COLUMNS: tuple[tuple[str, str, tuple[str, ...]], ...] = (
    ("vehicle_model", "Vehicle model", ("model", "vehicle", "car model", "car")),
    ("service", "Service", ("service name",)),
    ("part", "Part", ("part name", "item", "spare part")),
    ("quantity", "Qty", ("quantity",)),
    ("unit_price", "Unit price (excl. GST)", ("unit price", "price", "rate")),
    ("gst_percent", "GST %", ("gst", "gst percent", "gst rate", "tax %", "tax")),
    ("labour_hours", "Labour hours", ("labor hours", "labour", "labor")),
    ("duration_hours", "Duration hours", ("duration", "time hours")),
    ("keywords", "Keywords", ("keyword",)),
    ("includes", "Includes", ("includes services", "covers")),
)
EXPORT_HEADERS = tuple(header for _, header, _ in _COLUMNS)

_NON_ALPHANUMERIC = re.compile(r"[^0-9a-z%]+")
_LIST_SEPARATOR = re.compile(r"[;,\n]")
_MONEY_NOISE = re.compile(r"[₹,\s]|rs\.?|inr", re.IGNORECASE)


class PriceListFileError(ValueError):
    """The file cannot be read as a price list at all."""


@dataclass(frozen=True)
class ParsedPriceList:
    rows: tuple[PriceListRow, ...]
    # The spreadsheet row each parsed row came from (header is row 1 when
    # it is the first row).
    sheet_rows: tuple[int, ...]
    errors: tuple[PriceListIssue, ...]


def _header_key(text: str) -> str:
    return _NON_ALPHANUMERIC.sub(" ", text.casefold()).strip()


_HEADER_TO_FIELD: dict[str, str] = {}
for _field, _header, _aliases in _COLUMNS:
    for _name in (_header, _field.replace("_", " "), *_aliases):
        _HEADER_TO_FIELD.setdefault(_header_key(_name), _field)


def parse_price_list(filename: str, content: bytes) -> ParsedPriceList:
    if len(content) > MAX_UPLOAD_BYTES:
        raise PriceListFileError("The file is larger than 5 MB.")
    name = (filename or "").casefold()
    if name.endswith(".csv"):
        table = _read_csv(content)
    elif name.endswith(".xlsx") or name.endswith(".xlsm"):
        table = _read_xlsx(content)
    elif name.endswith(".xls"):
        raise PriceListFileError(
            "Old .xls files are not supported; save the sheet as .xlsx or .csv."
        )
    else:
        raise PriceListFileError("Upload an Excel (.xlsx) or CSV (.csv) file.")
    return _parse_table(table)


# A table is a list of (sheet row number, [(value, is_percent_format)]).
_Table = list[tuple[int, list[tuple[Any, bool]]]]


def _read_csv(content: bytes) -> _Table:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1252", errors="replace")
    reader = csv.reader(io.StringIO(text))
    return [
        (number, [(value, False) for value in row])
        for number, row in enumerate(reader, start=1)
        if number <= MAX_ROWS + _HEADER_SEARCH_ROWS + 1
    ]


def _read_xlsx(content: bytes) -> _Table:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise PriceListFileError("The file could not be opened as an Excel workbook.") from exc
    try:
        sheet = next(
            (ws for ws in workbook.worksheets if _header_key(ws.title) == "price list"),
            workbook.worksheets[0] if workbook.worksheets else None,
        )
        if sheet is None:
            raise PriceListFileError("The workbook has no sheets.")
        table: _Table = []
        for number, row in enumerate(sheet.iter_rows(), start=1):
            if number > MAX_ROWS + _HEADER_SEARCH_ROWS + 1:
                break
            table.append(
                (
                    number,
                    [
                        (cell.value, "%" in (getattr(cell, "number_format", "") or ""))
                        for cell in row
                    ],
                )
            )
        return table
    finally:
        workbook.close()


def _parse_table(table: _Table) -> ParsedPriceList:
    header_at = None
    columns: dict[int, str] = {}
    for index, (_, cells) in enumerate(table[:_HEADER_SEARCH_ROWS]):
        found = {
            position: _HEADER_TO_FIELD[_header_key(str(value))]
            for position, (value, _) in enumerate(cells)
            if value is not None and _header_key(str(value)) in _HEADER_TO_FIELD
        }
        if "service" in found.values():
            header_at, columns = index, found
            break
    if header_at is None:
        raise PriceListFileError(
            "No header row with a 'Service' column was found in the first "
            f"{_HEADER_SEARCH_ROWS} rows. Download the current price list to see the layout."
        )
    duplicates = {f for f in columns.values() if list(columns.values()).count(f) > 1}
    if duplicates:
        raise PriceListFileError(
            "These columns appear more than once: "
            + ", ".join(sorted(h for f, h, _ in _COLUMNS if f in duplicates))
        )

    rows: list[PriceListRow] = []
    sheet_rows: list[int] = []
    errors: list[PriceListIssue] = []
    for number, cells in table[header_at + 1 :]:
        values: dict[str, tuple[Any, bool]] = {}
        for position, field_name in columns.items():
            if position < len(cells):
                values[field_name] = cells[position]
        if all(_blank(value) for value, _ in values.values()):
            continue
        if len(rows) >= MAX_ROWS:
            errors.append(
                PriceListIssue("error", f"More than {MAX_ROWS} rows; the rest were not read.")
            )
            break

        def issue(message: str) -> None:
            errors.append(PriceListIssue("error", message, number))

        rows.append(
            PriceListRow(
                service=_text(values.get("service")) or "",
                vehicle_model=_text(values.get("vehicle_model")),
                part=_text(values.get("part")),
                quantity=_quantity(values.get("quantity"), issue),
                unit_price=_money(values.get("unit_price"), issue),
                gst_percent=_percent(values.get("gst_percent"), issue),
                labour_hours=_hours(values.get("labour_hours"), "Labour hours", issue),
                duration_hours=_hours(values.get("duration_hours"), "Duration hours", issue),
                keywords=_list(values.get("keywords")),
                includes=_list(values.get("includes")),
            )
        )
        sheet_rows.append(number)
    return ParsedPriceList(tuple(rows), tuple(sheet_rows), tuple(errors))


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _text(cell: tuple[Any, bool] | None) -> str | None:
    if cell is None or _blank(cell[0]):
        return None
    value = cell[0]
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def _list(cell: tuple[Any, bool] | None) -> tuple[str, ...]:
    text = _text(cell)
    if text is None:
        return ()
    return tuple(item.strip() for item in _LIST_SEPARATOR.split(text) if item.strip())


def _number(value: Any) -> Decimal | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            number = Decimal(str(value))
        except InvalidOperation:
            return None
    else:
        text = _MONEY_NOISE.sub("", str(value)).rstrip("%").strip()
        try:
            number = Decimal(text)
        except InvalidOperation:
            return None
    return number if number.is_finite() else None


def _money(cell: tuple[Any, bool] | None, issue) -> Decimal | None:
    if cell is None or _blank(cell[0]):
        return None
    number = _number(cell[0])
    if number is None:
        issue(f"Unit price {cell[0]!r} is not a number.")
    return number


def _percent(cell: tuple[Any, bool] | None, issue) -> Decimal | None:
    if cell is None or _blank(cell[0]):
        return None
    value, is_percent_format = cell
    number = _number(value)
    if number is None:
        issue(f"GST % {value!r} is not a number.")
        return None
    # A cell formatted as a percentage holds 18% as 0.18.
    if is_percent_format and not isinstance(value, str):
        number *= 100
    return Decimal(int(number)) if number == number.to_integral_value() else number


def _quantity(cell: tuple[Any, bool] | None, issue) -> int | None:
    if cell is None or _blank(cell[0]):
        return None
    number = _number(cell[0])
    if number is None or number != number.to_integral_value():
        issue(f"Qty {cell[0]!r} is not a whole number.")
        return None
    return int(number)


def _hours(cell: tuple[Any, bool] | None, label: str, issue) -> float | None:
    if cell is None or _blank(cell[0]):
        return None
    number = _number(cell[0])
    if number is None:
        issue(f"{label} {cell[0]!r} is not a number.")
        return None
    return float(number)


# --- Export -------------------------------------------------------------


def _export_values(row: PriceListRow) -> list[Any]:
    return [
        row.vehicle_model,
        row.service,
        row.part,
        row.quantity,
        row.unit_price,
        row.gst_percent,
        row.labour_hours,
        row.duration_hours,
        ", ".join(row.keywords) or None,
        ", ".join(row.includes) or None,
    ]


def export_csv(rows: tuple[PriceListRow, ...]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(EXPORT_HEADERS)
    for row in rows:
        writer.writerow(["" if v is None else v for v in _export_values(row)])
    # A BOM so Excel opens the UTF-8 file (₹, Tamil text) correctly.
    return buffer.getvalue().encode("utf-8-sig")


_INSTRUCTIONS = (
    "How to fill the price list",
    "",
    "One row per vehicle model + service + part. Prices are before GST.",
    "Vehicle model: leave blank for the price that applies to every model. A model's own "
    "rows are used for calls about that model; other models use the blank-model rows.",
    "Service: the service name, e.g. Brake Pad Replacement. Rows with the same service "
    "(and model) are one service.",
    "Part, Qty, Unit price: one part of the service. Qty defaults to 1. Leave Part blank "
    "on a row that only sets labour, duration or keywords.",
    "GST %: the part's GST rate. Blank uses the default GST % set on the Price List page.",
    "Labour hours / Duration hours: fill in on at least one row of each service and model. "
    "Duration defaults to the labour hours.",
    "Keywords: other words that mean the service in a call, separated by commas.",
    "Includes: services this one already includes (e.g. General Service includes Oil "
    "Change), separated by commas; they are charged once.",
    "The labour hourly rate, labour GST % and currency are set on the Price List page.",
)


def export_xlsx(rows: tuple[PriceListRow, ...]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Price list"
    sheet.append(list(EXPORT_HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for row in rows:
        sheet.append(_export_values(row))
    widths = (18, 26, 26, 6, 20, 8, 13, 15, 32, 26)
    for column, width in zip("ABCDEFGHIJ", widths):
        sheet.column_dimensions[column].width = width

    help_sheet = workbook.create_sheet("How to fill")
    for line in _INSTRUCTIONS:
        help_sheet.append([line])
    help_sheet["A1"].font = Font(bold=True)
    help_sheet.column_dimensions["A"].width = 120

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
